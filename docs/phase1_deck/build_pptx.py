"""Assemble the rendered slides into a 16:9 PowerPoint deck."""
import glob, os
from pptx import Presentation
from pptx.util import Inches

NOTES = {
    3:  "TO COMPLETE: replace the italic prompts with what you actually learned from the Dr. Agarwals Institute of Optometry clinical sessions. Keep only what is relevant to this project.",
    10: "TO COMPLETE: record the actual mentor feedback and what it changed in the product. Judges may verify this with your mentor - keep it factual.",
    1:  "TO COMPLETE: add your institution, mentor name and remaining team members.",
    8:  "Key slide: this is the Phase 0 to Phase 1 delta - the answer to 'what progress did you actually make?'",
    9:  "The measurement image is genuine pipeline output. Present it as the measurement stage working, not as a validated clinical result.",
}

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

pngs = sorted(glob.glob("render/slide*.png"))
for i, png in enumerate(pngs, start=1):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    if i in NOTES:
        s.notes_slide.notes_text_frame.text = NOTES[i]

out = "BeanHealth_CLR_Phase1_Deck.pptx"
prs.save(out)
print(f"wrote {out} - {len(pngs)} slides, {os.path.getsize(out)//1024} KB")
